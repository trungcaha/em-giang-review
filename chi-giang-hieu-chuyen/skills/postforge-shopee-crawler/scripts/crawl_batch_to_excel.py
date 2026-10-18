import argparse
import json
import sys
import time
from pathlib import Path

import httpx
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

BASE_URL = "http://127.0.0.1:8765"
IMAGES_BASE_DIR = Path(r"D:\video-ai-aff\shopee-aff-crawler\storage\images")


def save_to_excel(rows, out_path: Path, sub_prefix: str = "PostForge"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Shopee Affiliate"

    headers = [
        "STT",
        "Nhóm sản phẩm",
        "Tên sản phẩm yêu cầu",
        "Tên sản phẩm Shopee",
        "Giá",
        "Hoa hồng",
        "Link Shopee Affiliate",
        "Link Affiliate kèm Sub_id (Theo Kênh)",
        "Đường dẫn Folder Ảnh",
        "Số ảnh",
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, r in enumerate(rows, start=1):
        base_link = (r.get("aff_link") or "").strip()
        code = base_link.rstrip("/").split("/")[-1] if base_link else f"SP{idx:02d}"
        ch = r.get("sub_channel") or sub_prefix
        sub_link = f"{base_link}?sub_id={ch}-Tap01-FBReels-{code}" if base_link else ""

        row_data = [
            idx,
            r.get("group", ""),
            r.get("req_name", ""),
            r.get("title", ""),
            r.get("price_text", ""),
            r.get("commission_text", ""),
            base_link,
            sub_link,
            r.get("folder_path", ""),
            r.get("image_count", 0),
        ]
        ws.append(row_data)
        row_num = idx + 1
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=row_num, column=col_idx)
            c.font = Font(name="Arial", size=10)
            c.border = thin_border
            c.alignment = Alignment(
                horizontal="center" if col_idx in (1, 5, 6, 10) else "left",
                vertical="center",
                wrap_text=True,
            )

    col_widths = [6, 22, 34, 48, 14, 12, 32, 54, 62, 10]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 28

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"  [Excel] Đã lưu {len(rows)} sản phẩm -> {out_path}")


def crawl_single(client: httpx.Client, keyword: str, must_contain: str, seen_ids: set):
    r = client.post(
        f"{BASE_URL}/api/crawl",
        json={"keyword": keyword, "max_items": 1, "must_contain": must_contain},
    )
    if r.status_code != 200:
        print("  ! Lỗi gọi /api/crawl:", r.text)
        return None
    time.sleep(2)
    while True:
        st = client.get(f"{BASE_URL}/api/status").json()
        print(f"    -> {st.get('crawler_status')}")
        if not st.get("crawler_running"):
            break
        time.sleep(3)

    items = client.get(f"{BASE_URL}/api/products").json().get("items", [])
    for it in items:
        if it["item_id"] not in seen_ids and it.get("keyword", "").lower() == keyword.lower():
            return it
    for it in items:
        if it.get("keyword", "").lower() == keyword.lower():
            return it
    return None


def main():
    parser = argparse.ArgumentParser(description="Batch crawl Shopee Affiliate products via PostForge")
    parser.add_argument("--input-json", required=True, help="Path to JSON file containing list of products to crawl")
    parser.add_argument("--output-excel", required=True, help="Path to output .xlsx file")
    parser.add_argument("--sub-prefix", default="PostForge", help="Default Sub_id channel prefix")
    args = parser.parse_args()

    product_list = json.loads(Path(args.input_json).read_text(encoding="utf-8"))
    out_excel = Path(args.output_excel)

    client = httpx.Client(timeout=30.0)
    st = client.get(f"{BASE_URL}/api/status").json()
    if not st.get("adb_connected") and not st.get("maxcloud_connected"):
        print("Chưa kết nối ADB điện thoại trên PostForge!")
        sys.exit(1)

    existing_items = client.get(f"{BASE_URL}/api/products").json().get("items", [])
    seen_ids = {x["item_id"] for x in existing_items}
    collected_rows = []

    for i, p in enumerate(product_list, start=1):
        req_name = p.get("req_name") or p.get("keyword")
        keyword = p.get("keyword") or req_name
        must_contain = p.get("must_contain", "")
        fallback_kw = p.get("fallback_keyword") or keyword
        group = p.get("group", "")
        ref_comm = p.get("ref_comm", "")

        print(f"\n[{i}/{len(product_list)}] ({group}) {req_name}")
        found = crawl_single(client, keyword, must_contain, seen_ids)
        if not found and (must_contain or fallback_kw != keyword):
            print(f"  -> Thử lại fallback keyword: {fallback_kw}")
            found = crawl_single(client, fallback_kw, "", seen_ids)

        if found:
            seen_ids.add(found["item_id"])
            comm_text = found.get("commission_text") or (f"HH {ref_comm}" if ref_comm else "HH —")
            collected_rows.append(
                {
                    "group": group,
                    "req_name": req_name,
                    "item_id": found["item_id"],
                    "title": found["title"],
                    "price_text": found.get("price_text", ""),
                    "commission_text": comm_text,
                    "aff_link": found.get("aff_link", ""),
                    "folder_path": str(IMAGES_BASE_DIR / str(found["item_id"])),
                    "image_count": found.get("image_count", 0),
                    "sub_channel": p.get("sub_channel", args.sub_prefix),
                }
            )
        else:
            collected_rows.append(
                {
                    "group": group,
                    "req_name": req_name,
                    "item_id": "",
                    "title": "(Không tìm thấy trên Shopee)",
                    "price_text": "",
                    "commission_text": "",
                    "aff_link": "",
                    "folder_path": "",
                    "image_count": 0,
                    "sub_channel": p.get("sub_channel", args.sub_prefix),
                }
            )
        save_to_excel(collected_rows, out_excel, sub_prefix=args.sub_prefix)


if __name__ == "__main__":
    main()
