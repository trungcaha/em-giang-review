# -*- coding: utf-8 -*-
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.comments import Comment
import data as D

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="0F4C5C")
SUB_FILL = PatternFill("solid", fgColor="E6F0F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF8D6")
WARN_FILL = PatternFill("solid", fgColor="FDE2E1")
THIN = Side(style="thin", color="D0D7DE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

F_BASE = Font(name=FONT, size=10)
F_HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=16, bold=True, color="0F4C5C")
F_SUB = Font(name=FONT, size=10, italic=True, color="555555")
F_BOLD = Font(name=FONT, size=10, bold=True)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")      # liên kết sang sheet khác
F_URL = Font(name=FONT, size=10, color="0563C1", underline="single")
F_SECTION = Font(name=FONT, size=12, bold=True, color="0F4C5C")

wb = Workbook()
wb.remove(wb.active)
PRODUCTS = D.load_products()
import math
_A = [p["per_order"] for p in PRODUCTS if p["tier"] == "A"] or [p["per_order"] for p in PRODUCTS]
AVG_A = sum(_A) / len(_A)
HEADS = sum(math.ceil(D.CAPACITY["target"] / D.CAPACITY[r["cap"]]) if r["cap"] else int(r["count"]) for r in D.ROLES)
_F = D.FORECAST
MAXD = D.MAX_DAY
MONTHLY = MAXD * 30 * _F["views"] * _F["ctr"] * _F["cr"] * AVG_A
COST = HEADS * _F["cost"]
BREAKEVEN = COST / (MAXD * 30 * _F["ctr"] * _F["cr"] * AVG_A)
vn = lambda x: f"{x:,.0f}".replace(",", ".")
P = {p["stt"]: p for p in PRODUCTS}
NP = len(PRODUCTS)
P_FIRST, P_LAST = 5, 5 + NP - 1
KB_FIRST, KB_LAST = 5, 404
TD_FIRST = 5
DAYS = (D.WEEKS[-1]["end"] - D.WEEKS[0]["start"]).days + 1
TD_LAST = TD_FIRST + DAYS - 1
CV_FIRST, CV_LAST = 5, 60
NPIL = len(D.PILLARS)
TOP = sorted(PRODUCTS, key=lambda p: -p["per_order"])[:5]
TOP_TXT = ", ".join(f"{p['name']} ~{round(p['per_order']/1000)}k" for p in TOP)


def sheet(name, title, subtitle, widths):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    ws["A2"] = subtitle
    ws["A2"].font = F_SUB
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    return ws


def header(ws, row, labels, col=1):
    for i, h in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=h)
        c.font, c.fill, c.alignment, c.border = F_HDR, HDR_FILL, CENTER, BORDER
    ws.row_dimensions[row].height = 30


def put(ws, row, col, value, font=F_BASE, fill=None, fmt=None, align=WRAP):
    c = ws.cell(row=row, column=col, value=value)
    c.font, c.border, c.alignment = font, BORDER, align
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    return c


def section(ws, row, text):
    ws.cell(row=row, column=1, value=text).font = F_SECTION


def dv_list(ws, items, rng):
    dv = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)


# =================================================================== HUONG_DAN
ws = sheet("HUONG_DAN", "HƯỚNG DẪN – FILE VẬN HÀNH " + D.PROJECT["name"].upper(),
           D.PROJECT["subtitle"] + " · lập ngày " + D.PROJECT["today"].strftime("%d/%m/%Y"), [4, 30, 92])
r = 4
section(ws, r, "Mục tiêu"); r += 1
for k, v in [("Mục tiêu sản lượng", D.PROJECT["goal"]),
             ("Chỉ số chính", D.PROJECT["north_star"] + " (xem TONG_QUAN)"),
             ("Bắt đầu", "Giai đoạn 0 dựng nền 16–18/10 · Tuần 1 bắt đầu thứ Hai 19/10/2026 · lộ trình 12 tuần đến 10/01/2027")]:
    put(ws, r, 2, k, F_BOLD); put(ws, r, 3, v); r += 1
r += 1
section(ws, r, "Bốn điều cần biết trước khi chạy"); r += 1
for k, v in [("1. Hoa hồng tính theo Shopee, không theo TikTok",
              "Tên, giá, hoa hồng ở SAN_PHAM lấy đúng sản phẩm trong link Shopee. Thực phẩm chức năng trên Shopee thường chỉ 2,5–7,5% (TikTok 10–15%), nên thứ hạng đã đổi. Hoa hồng/đơn cao nhất hiện tại: " + TOP_TXT + "."),
             ("2. Mỗi video một Sub_id",
              f"Cột I của KHO_KICH_BAN tự tạo link dạng …?sub_id={D.SUB_PREFIX}-<ID video>-FBReels-SP<mã SP>. Dán đúng link này vào bình luận gốc thì báo cáo Shopee cho biết video nào ra đơn – không cần đoán."),
             ("3. Một page, nội dung nhạy cảm",
              "Chị Giang Hiểu Chuyện là page duy nhất của dự án: bị hạn chế là mất hết. Giữ giọng chia sẻ kiến thức, không xưng bác sĩ, không hình gợi dục, không hứa “tăng sinh lý” (QC điểm 3–5)."),
             ("4. Bài toán kinh tế chưa được chứng minh",
              f"Với giả định {vn(_F['views'])} view/video, hoa hồng khi đủ {MAXD} video/ngày khoảng {MONTHLY/1e6:.1f} triệu/tháng, trong khi {HEADS} người tốn khoảng {COST/1e6:.0f} triệu/tháng. Ngưỡng hòa vốn ≈ {vn(BREAKEVEN)} view/video. Giai đoạn 1 chạy 2 người kiêm nhiệm; lấy số thật của page ngay (việc T08).")]:
    put(ws, r, 2, k, F_BOLD, WARN_FILL); put(ws, r, 3, v); ws.row_dimensions[r].height = 56; r += 1
r += 1
section(ws, r, "Quy ước màu"); r += 1
for k, v, f, fl in [("Chữ xanh dương, nền vàng nhạt", "Ô cần điền hoặc chọn – chỉ sửa những ô này", F_INPUT, INPUT_FILL),
                    ("Chữ đen", "Công thức – không sửa, tự tính", F_BASE, None),
                    ("Chữ xanh lá", "Công thức lấy dữ liệu từ sheet khác", F_LINK, None),
                    ("Nền đỏ nhạt", "Cần xử lý (quá hạn, dưới ngưỡng, link cần sửa)", F_BASE, WARN_FILL)]:
    put(ws, r, 2, k, f, fl); put(ws, r, 3, v); r += 1
r += 1
section(ws, r, "Ai dùng sheet nào"); r += 1
for k, v in [("Điều phối kiêm trưởng nội dung (R1)", "TONG_QUAN mỗi sáng · KHO_KICH_BAN (duyệt: đổi trạng thái) · LO_TRINH · CONG_VIEC · DU_BAO · review Chủ nhật"),
             ("Biên kịch – master prompt (R2)", "KHO_KICH_BAN: nhận dòng “Ý tưởng” → viết prompt → đổi sang “Chờ duyệt”. Khuôn và luật giọng nói ở KENH, mẫu ở TAI_NGUYEN"),
             ("Sản xuất video (R3)", "KHO_KICH_BAN: lọc “Đã duyệt” → Viral Studio → “Đang sản xuất” → “QC”"),
             ("Đăng bài, cộng đồng & dữ liệu (R4)", "KHO_KICH_BAN: lọc “QC” → kiểm 10 điểm (QUY_TRINH) → Postiz, dán link cột I vào bình luận gốc → “Đã đăng”. SAN_PHAM (trạng thái link) · THEO_DOI_NGAY (nhập số mỗi sáng)"),
             ("Mọi người", "NGUON_CONTENT: thấy video, kênh, hook hay thì dán link + ghi “Học được gì”. Trang HTML (tab Nguồn content) dán link là xem ngay")]:
    put(ws, r, 2, k, F_BOLD); put(ws, r, 3, v); ws.row_dimensions[r].height = 30; r += 1
r += 1
section(ws, r, "Luồng một video"); r += 1
put(ws, r, 2, "Trạng thái", F_BOLD); put(ws, r, 3, " → ".join(D.STATUSES[:-1]) + " (hoặc Hủy)"); r += 1
put(ws, r, 2, "Cập nhật file", F_BOLD); put(ws, r, 3, "File HTML “du-an-chi-giang.html” là bản trực quan để xem và giao việc; số liệu sống nằm ở file Excel này. Chạy python build.py chỉ làm mới file HTML; file Excel chỉ tạo lại khi chạy python build.py excel (sẽ xóa số liệu đã nhập).")

# =================================================================== KENH
C = D.CHANNEL
ws = sheet("KENH", "KÊNH CHỊ GIANG HIỂU CHUYỆN – HỒ SƠ & 5 TRỤ CỘT", "Một kênh, chia nội dung thành 5 trụ cột. Số video/ngày của từng trụ cột lấy từ LO_TRINH.",
           [8, 26, 44, 46, 46, 14, 20])
r = 4
section(ws, r, "Hồ sơ kênh"); r += 1
for k, v in [("Tên page", C["page"]), ("Postiz ID", C["postiz_id"]), ("Thư mục video", C["folder"]), ("Trạng thái", C["status"]),
             ("Nhân vật", C["host"]), ("Định dạng", C["format"]), ("Công cụ", C["tool"])]:
    put(ws, r, 2, k, F_BOLD); put(ws, r, 3, v); ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
    ws.row_dimensions[r].height = 30 if len(v) > 90 else 16; r += 1
r += 1
section(ws, r, "Tệp khán giả (Facebook Insights lifetime, 16/10/2026)"); r += 1
for k, v in C["audience"]:
    put(ws, r, 2, k, F_BOLD); put(ws, r, 3, v); ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7); r += 1
r += 1
section(ws, r, "Luật giọng nói (áp vào mọi prompt)"); r += 1
for v in C["voice"]:
    put(ws, r, 2, "•", align=CENTER); put(ws, r, 3, v); ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7); r += 1
r += 1
section(ws, r, "5 trụ cột nội dung"); r += 1
header(ws, r, ["Mã", "Trụ cột", "Nội dung", "Câu hỏi mẫu", "Sản phẩm chủ lực (mã SP)", "Video/ngày (tuần 12)", "Phụ trách viết"])
r0 = r + 1
for i, pl in enumerate(D.PILLARS):
    r = r0 + i
    put(ws, r, 1, pl["code"], F_BOLD, align=CENTER)
    put(ws, r, 2, pl["name"], F_BOLD); put(ws, r, 3, pl["desc"])
    put(ws, r, 4, "\n".join("• " + x for x in pl["examples"]))
    put(ws, r, 5, "\n".join(f"{n}. {P[n]['name']}" for n in pl["products"] if n in P))
    put(ws, r, 6, f"=LO_TRINH!{get_column_letter(4 + i)}15", F_LINK, align=CENTER)
    put(ws, r, 7, None, F_INPUT, INPUT_FILL)
    ws.row_dimensions[r].height = 80
put(ws, r0 + NPIL, 5, "Tổng", F_BOLD); put(ws, r0 + NPIL, 6, f"=SUM(F{r0}:F{r0+NPIL-1})", F_BOLD, align=CENTER)

# =================================================================== LO_TRINH
ws = sheet("LO_TRINH", "LỘ TRÌNH 12 TUẦN – VIDEO/NGÀY THEO TRỤ CỘT",
           "Ô xanh dương = số video/ngày của từng trụ cột, chỉnh được. THEO_DOI_NGAY, DU_BAO, KENH tự lấy số từ bảng này.",
           [9, 12, 12, 8, 8, 8, 8, 8, 11, 12, 50])
header(ws, 3, ["Tuần", "Từ ngày", "Đến ngày"] + D.PILLAR_CODES + ["Tổng/ngày", "Tổng/tuần", "Mốc / sự kiện"])
for i, w in enumerate(D.WEEKS):
    r = 4 + i
    put(ws, r, 1, w["week"], F_BOLD, align=CENTER)
    put(ws, r, 2, w["start"], fmt="dd/mm/yyyy", align=CENTER)
    put(ws, r, 3, f"=B{r}+6", fmt="dd/mm/yyyy", align=CENTER)
    for j, k in enumerate(D.PILLAR_CODES):
        put(ws, r, 4 + j, w[k], F_INPUT, INPUT_FILL, align=CENTER)
    put(ws, r, 9, f"=SUM(D{r}:H{r})", F_BOLD, align=CENTER)
    put(ws, r, 10, f"=I{r}*7", fmt="#,##0", align=CENTER)
    put(ws, r, 11, w["event"])
put(ws, 16, 1, "Tổng 12 tuần", F_BOLD)
put(ws, 16, 10, "=SUM(J4:J15)", F_BOLD, fmt="#,##0", align=CENTER)
ws.conditional_formatting.add("D4:H15", CellIsRule(operator="equal", formula=["0"], fill=PatternFill("solid", fgColor="EEEEEE")))
ws.conditional_formatting.add("I4:I15", CellIsRule(operator="greaterThanOrEqual", formula=[str(MAXD)], fill=PatternFill("solid", fgColor="C8E6C9")))
put(ws, 17, 1, " · ".join(f"{p['code']} {p['name']}" for p in D.PILLARS), F_SUB)
ws.merge_cells(start_row=17, start_column=1, end_row=17, end_column=11)
section(ws, 19, "Các giai đoạn và điều kiện qua cửa")
header(ws, 20, ["Giai đoạn", "", "Thời gian", "", "", "Mục tiêu", "", "", "", "", "Điều kiện để sang giai đoạn sau"])
for i, ph in enumerate(D.PHASES):
    r = 21 + i
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=10)
    put(ws, r, 1, ph["phase"], F_BOLD); put(ws, r, 3, ph["time"]); put(ws, r, 6, ph["goal"]); put(ws, r, 11, ph["gate"])
    ws.row_dimensions[r].height = 45
ws.merge_cells(start_row=20, start_column=1, end_row=20, end_column=2)
ws.merge_cells(start_row=20, start_column=3, end_row=20, end_column=5)
ws.merge_cells(start_row=20, start_column=6, end_row=20, end_column=10)
ws.freeze_panes = "B4"

# =================================================================== NHAN_SU
ws = sheet("NHAN_SU", "NHÂN SỰ & PHÂN VAI", "Số người cần tự tính theo công suất (ô xanh dương). Điền tên người vào cột “Họ tên”.",
           [32, 12, 22, 16, 50, 32, 34, 36, 10])
section(ws, 3, "Thông số công suất")
for r, (k, v, f, fl, fmt) in zip(range(4, 8), [
        ("Video mục tiêu/ngày (lấy từ LO_TRINH)", "=MAX(LO_TRINH!I4:I15)", F_LINK, None, "0"),
        ("Prompt 1 biên kịch viết/ngày", D.CAPACITY["kb"], F_INPUT, INPUT_FILL, "0"),
        ("Video 1 người sản xuất/ngày", D.CAPACITY["dung"], F_INPUT, INPUT_FILL, "0"),
        ("Bài 1 người QC & đăng/ngày", D.CAPACITY["qc"], F_INPUT, INPUT_FILL, "0")]):
    put(ws, r, 1, k, F_BOLD); put(ws, r, 2, v, f, fl, fmt, CENTER)
ws["B5"].comment = Comment("Giả định: prompt sinh từ khuôn có sẵn (_tao-prompt-*.py), ~20 phút/prompt kể cả sửa.", "Điều phối")
ws["B6"].comment = Comment("Giả định: Viral Studio chạy hàng đợi, người sản xuất chỉ soát shot và làm lại shot lỗi. Chỉnh theo tuần 1.", "Điều phối")
header(ws, 10, ["Vai trò", "Số người cần", "Họ tên (điền)", "Kiêm vai", "Trách nhiệm chính", "Đầu ra mỗi ngày", "KPI", "Công cụ / skill", "Báo cáo cho"])
for i, ro in enumerate(D.ROLES):
    r = 11 + i
    put(ws, r, 1, f"{ro['code']} · {ro['role']}", F_BOLD)
    cnt = ro["count"]
    if cnt.startswith("="):
        cnt = cnt.replace("B4", "$B$4").replace("B5", "$B$5").replace("B6", "$B$6").replace("B7", "$B$7")
        put(ws, r, 2, cnt, align=CENTER, fmt="0")
    else:
        put(ws, r, 2, int(cnt), F_INPUT, INPUT_FILL, "0", CENTER)
    put(ws, r, 3, None, F_INPUT, INPUT_FILL)
    put(ws, r, 4, None, F_INPUT, INPUT_FILL)
    for col, v in [(5, ro["duty"]), (6, ro["output"]), (7, ro["kpi"]), (8, ro["tools"]), (9, ro["report"])]:
        put(ws, r, col, v)
    ws.row_dimensions[r].height = 62
last = 11 + len(D.ROLES)
put(ws, last, 1, f"TỔNG NHÂN SỰ KHI ĐỦ {MAXD} VIDEO/NGÀY", F_BOLD, SUB_FILL)
put(ws, last, 2, f"=SUM(B11:B{last-1})", F_BOLD, SUB_FILL, "0", CENTER)
put(ws, last + 1, 1, "Gợi ý giai đoạn 1 (5–6 video/ngày): 2 người – người 1 kiêm R1 + R2 (chọn chủ đề, viết và duyệt prompt); người 2 kiêm R3 + R4 (Viral Studio, QC, đăng, bình luận, số liệu). Tăng người khi view thực tế vượt ngưỡng hòa vốn ở DU_BAO.", F_SUB)
ws.merge_cells(start_row=last + 1, start_column=1, end_row=last + 1, end_column=9)
ws.row_dimensions[last + 1].height = 30
ws.freeze_panes = "B11"

# =================================================================== QUY_TRINH
ws = sheet("QUY_TRINH", "QUY TRÌNH VẬN HÀNH (SOP)", "10 bước cho một video, từ ý tưởng đến số liệu.",
           [6, 28, 18, 10, 12, 36, 46, 6, 6, 6])
header(ws, 4, ["Bước", "Việc", "Khi nào", "Làm (R)", "Phối hợp", "Đầu vào", "Đầu ra"])
for i, s in enumerate(D.SOP):
    r = 5 + i
    for col, v in enumerate(s, 1):
        put(ws, r, col, v, F_BOLD if col == 2 else F_BASE, align=CENTER if col in (1, 4) else WRAP)
    ws.row_dimensions[r].height = 36
r = 5 + len(D.SOP) + 1
section(ws, r, "Ma trận trách nhiệm RACI (R = làm, A = chịu trách nhiệm cuối, C = tham vấn, I = được báo)"); r += 1
header(ws, r, ["Bước", "Việc"] + D.RACI_ROLES)
raci_fill = {"R": PatternFill("solid", fgColor="BBDEFB"), "A": PatternFill("solid", fgColor="FFCC80"),
             "C": PatternFill("solid", fgColor="E1F5FE"), "I": PatternFill("solid", fgColor="F5F5F5")}
for s in D.SOP:
    r += 1
    put(ws, r, 1, s[0], align=CENTER); put(ws, r, 2, s[1])
    for j, role in enumerate(D.RACI_ROLES):
        v = D.RACI[s[0]].get(role, "")
        put(ws, r, 3 + j, v, F_BOLD, raci_fill.get(v), align=CENTER)
r += 1
put(ws, r, 2, " · ".join(f"{ro['code']} {ro['role']}" for ro in D.ROLES), F_SUB)
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
r += 2
section(ws, r, "Checklist QC 10 điểm (đạt hết mới được đăng)"); r += 1
header(ws, r, ["#", "Tiêu chí", "", "", "", "", ""])
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
for i, q in enumerate(D.QC_CHECK, 1):
    r += 1
    put(ws, r, 1, i, align=CENTER); put(ws, r, 2, q)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
r += 2
section(ws, r, "Nhịp làm việc một ngày"); r += 1
header(ws, r, ["Giờ", "Việc", "", "Ai", "", "Ghi chú", ""])
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
for t in D.DAILY_RHYTHM:
    r += 1
    put(ws, r, 1, t[0], F_BOLD); put(ws, r, 2, t[1]); put(ws, r, 4, t[2]); put(ws, r, 6, t[3])
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
r += 2
section(ws, r, "Họp định kỳ"); r += 1
header(ws, r, ["", "Khi nào", "Nội dung", "", "", "", "Chủ trì"])
for t in D.RITUALS:
    r += 1
    put(ws, r, 2, t[0], F_BOLD); put(ws, r, 3, t[1]); put(ws, r, 7, t[2])
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
r += 2
section(ws, r, "Quy tắc đặt tên file video"); r += 1
put(ws, r, 2, "CG-<YYYYMMDD>-<giờ đăng>-<ID video>-<slug>.mp4 → ví dụ CG-20261020-0730-T1-002-review-vuong-bao.mp4 · lưu vào viral-studio-downloads/videos/gioi-tinh/", F_BASE)
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)

# =================================================================== LICH_DANG
ws = sheet("LICH_DANG", f"LỊCH ĐĂNG – {len(D.SLOTS)} KHUNG GIỜ/NGÀY", "Giờ theo thói quen lướt Facebook của đàn ông 45+: cà phê sáng, nghỉ trưa, tối muộn. Tuần 1–3 chưa đủ 8 video thì bỏ bớt khung 13:00 và 22:45 trước.",
           [12, 42, 32, 26])
header(ws, 4, ["Khung giờ", "Loại video", "Mục tiêu", "Gắn sản phẩm"])
for i, s in enumerate(D.SLOTS):
    for col, v in enumerate(s, 1):
        put(ws, 5 + i, col, v, F_BOLD if col == 1 else F_BASE)
r = 5 + len(D.SLOTS) + 1
section(ws, r, "Chủ đề theo thứ trong tuần"); r += 1
header(ws, r, ["Thứ", "Chủ đề ngày", "Bối cảnh gợi ý", "Sản phẩm tập trung"])
for t in D.WEEKDAY_THEMES:
    r += 1
    for col, v in enumerate(t, 1):
        put(ws, r, col, v, F_BOLD if col == 1 else F_BASE)
r += 2
section(ws, r, "Lịch mùa vụ 10/2026 – 01/2027"); r += 1
header(ws, r, ["Thời điểm", "Sự kiện", "Hướng nội dung", ""])
ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
for t in D.SEASON:
    r += 1
    put(ws, r, 1, t[0], F_BOLD); put(ws, r, 2, t[1]); put(ws, r, 3, t[2])
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
r += 2
section(ws, r, "Quy tắc chống lặp"); r += 1
for t in ["Một sản phẩm tối đa 2 lần/ngày, mỗi lần đổi bối cảnh và góc kể",
          "Không dùng chung câu mở đầu quá 3 lần/tuần; băng chữ đỏ mỗi video một câu hỏi khác",
          "Biến thể của video thắng cách nhau ≥ 5 ngày",
          "Postiz giới hạn ~30 request/giờ: 8 bài × 2 request, script tự chờ"]:
    put(ws, r, 1, "•"); put(ws, r, 2, t); ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4); r += 1

# =================================================================== SAN_PHAM
ws = sheet("SAN_PHAM", f"SẢN PHẨM AFFILIATE SHOPEE – {NP} LINK",
           "Tham chiếu duy nhất: sản phẩm trong link Shopee (tên, giá, hoa hồng). Hạng xếp theo hoa hồng/đơn trên Shopee và độ hợp chủ đề. Link kèm Sub_id ở cột M là link chung; link riêng từng video nằm ở KHO_KICH_BAN cột I.",
           [6, 20, 30, 44, 11, 8, 11, 9, 9, 7, 40, 28, 46, 40, 7, 30, 16, 10, 10])
header(ws, 4, ["Mã SP", "Nhóm", "Tên sản phẩm (theo link Shopee)", "Tên đầy đủ trên Shopee", "Giá (đ)", "HH %", "HH/đơn (đ)",
               "Trụ cột chính", "Trụ cột phụ", "Hạng", "Góc nội dung", "Link Shopee Affiliate", "Link kèm Sub_id (chung)", "Thư mục ảnh", "Số ảnh",
               "Ghi chú", "Trạng thái link", "Video trong kho", "Đơn ghi nhận"])
tier_fill = {"A": PatternFill("solid", fgColor="C8E6C9"), "B": PatternFill("solid", fgColor="FFF59D"), "C": PatternFill("solid", fgColor="EEEEEE")}
for i, p in enumerate(PRODUCTS):
    r = P_FIRST + i
    put(ws, r, 1, p["stt"], F_BOLD, align=CENTER)
    put(ws, r, 2, p["group"]); put(ws, r, 3, p["name"], F_BOLD); put(ws, r, 4, p["shopee"])
    put(ws, r, 5, p["price"], fmt="#,##0", align=CENTER)
    put(ws, r, 6, p["hh"], fmt="0.0%", align=CENTER)
    put(ws, r, 7, f"=E{r}*F{r}", fmt="#,##0", align=CENTER)
    put(ws, r, 8, p["ch"], F_BOLD, align=CENTER); put(ws, r, 9, p["ch2"], align=CENTER)
    put(ws, r, 10, p["tier"], F_BOLD, tier_fill[p["tier"]], align=CENTER)
    put(ws, r, 11, p["angle"])
    c = put(ws, r, 12, p["link"], F_URL); c.hyperlink = p["link"]
    c = put(ws, r, 13, p["sublink"], F_URL); c.hyperlink = p["sublink"]
    c = put(ws, r, 14, p["folder"], F_URL)
    if p["folder"]:
        c.hyperlink = "file:///" + str(p["folder"]).replace("\\", "/")
    put(ws, r, 15, p["images"], align=CENTER)
    put(ws, r, 16, p["note"] or None)
    put(ws, r, 17, "Chưa kiểm tra", F_INPUT, INPUT_FILL, align=CENTER)
    put(ws, r, 18, f"=COUNTIF(KHO_KICH_BAN!$G${KB_FIRST}:$G${KB_LAST},A{r})", F_LINK, align=CENTER)
    put(ws, r, 19, f"=SUMIF(KHO_KICH_BAN!$G${KB_FIRST}:$G${KB_LAST},A{r},KHO_KICH_BAN!$S${KB_FIRST}:$S${KB_LAST})", F_LINK, align=CENTER)
    ws.row_dimensions[r].height = 42
dv_list(ws, ["Chưa kiểm tra", "OK", "Cần sửa", "Đã thay link"], f"Q{P_FIRST}:Q{P_LAST}")
dv_list(ws, ["A", "B", "C"], f"J{P_FIRST}:J{P_LAST}")
dv_list(ws, D.PILLAR_CODES, f"H{P_FIRST}:I{P_LAST}")
ws.conditional_formatting.add(f"Q{P_FIRST}:Q{P_LAST}", CellIsRule(operator="equal", formula=['"Cần sửa"'], fill=WARN_FILL))
ws.conditional_formatting.add(f"Q{P_FIRST}:Q{P_LAST}", CellIsRule(operator="equal", formula=['"OK"'], fill=PatternFill("solid", fgColor="C8E6C9")))
ws.auto_filter.ref = f"A4:S{P_LAST}"
ws.freeze_panes = "D5"

# =================================================================== KHO_KICH_BAN
ws = sheet("KHO_KICH_BAN", "KHO KỊCH BẢN & THEO DÕI TỪNG VIDEO",
           f"Mỗi dòng = 1 video. Nhập Mã SP → tên SP, link kèm Sub_id riêng của video, thư mục ảnh tự hiện. Dán link cột I vào bình luận gốc. Dòng 5 là dòng ví dụ.",
           [10, 8, 18, 40, 40, 26, 7, 34, 44, 26, 14, 14, 14, 11, 9, 24, 9, 8, 7, 24])
header(ws, 4, ["ID", "Trụ cột", "Loại video", "Tiêu đề / câu hỏi", "Hook 3 giây", "Bối cảnh", "Mã SP", "Tên SP (tự động)",
               "Link kèm Sub_id của video (tự động)", "Thư mục ảnh (tự động)", "Biên kịch", "Người sản xuất", "Trạng thái", "Ngày đăng",
               "Khung giờ", "Link bài đăng", "View", "Click", "Đơn", "Ghi chú"])
cnt = {}
for r in range(KB_FIRST, KB_LAST + 1):
    i = r - KB_FIRST
    idea = D.IDEAS[i] if i < len(D.IDEAS) else None
    if idea:
        pil, typ, title, hook, scene, sp = idea
        cnt[pil] = cnt.get(pil, 0) + 1
        vals = [f"{pil}-{cnt[pil]:03d}", pil, typ, title, hook, scene, sp if sp in P else None]
    else:
        vals = [None] * 7
    for col, v in enumerate(vals, 1):
        put(ws, r, col, v, F_INPUT, INPUT_FILL, align=CENTER if col in (1, 2, 7) else WRAP)
    m = f"MATCH($G{r},SAN_PHAM!$A${P_FIRST}:$A${P_LAST},0)"
    put(ws, r, 8, f'=IF($G{r}="","",IFERROR(INDEX(SAN_PHAM!$D${P_FIRST}:$D${P_LAST},{m}),"Sai mã SP"))', F_LINK)
    put(ws, r, 9, f'=IF(OR($G{r}="",$A{r}=""),"",IFERROR(INDEX(SAN_PHAM!$L${P_FIRST}:$L${P_LAST},{m})&"?sub_id={D.SUB_PREFIX}-"&SUBSTITUTE($A{r},"-","")&"-FBReels-SP"&TEXT($G{r},"00"),""))', F_LINK)
    put(ws, r, 10, f'=IF($G{r}="","",IFERROR(INDEX(SAN_PHAM!$N${P_FIRST}:$N${P_LAST},{m}),""))', F_LINK)
    for col in range(11, 21):
        put(ws, r, col, "Ý tưởng" if (col == 13 and idea) else None, F_INPUT, INPUT_FILL,
            fmt="dd/mm/yyyy" if col == 14 else ("#,##0" if col in (17, 18, 19) else None),
            align=CENTER if col in (13, 14, 15, 17, 18, 19) else WRAP)
# dòng ví dụ
for col, v in [(11, "Lan (ví dụ)"), (12, "Minh (ví dụ)"), (13, "Đã đăng"), (14, dt.date(2026, 10, 19)), (15, "11:30"),
               (16, "https://facebook.com/reel/… (ví dụ)"), (17, 12000), (18, 60), (19, 2), (20, "DÒNG VÍ DỤ – số liệu minh họa, xóa khi dùng thật")]:
    ws.cell(row=KB_FIRST, column=col, value=v)
dv_list(ws, D.STATUSES, f"M{KB_FIRST}:M{KB_LAST}")
dv_list(ws, D.VIDEO_TYPES, f"C{KB_FIRST}:C{KB_LAST}")
dv_list(ws, D.PILLAR_CODES, f"B{KB_FIRST}:B{KB_LAST}")
dv_list(ws, [s[0] for s in D.SLOTS], f"O{KB_FIRST}:O{KB_LAST}")
status_colors = {"Chờ duyệt": "FFF59D", "Cần sửa": "FFCDD2", "Đã duyệt": "BBDEFB", "Đang sản xuất": "D1C4E9", "QC": "FFE0B2",
                 "Sẵn sàng đăng": "B2DFDB", "Đã đăng": "C8E6C9", "Hủy": "E0E0E0"}
for s, col in status_colors.items():
    ws.conditional_formatting.add(f"M{KB_FIRST}:M{KB_LAST}", CellIsRule(operator="equal", formula=[f'"{s}"'], fill=PatternFill("solid", fgColor=col)))
ws.auto_filter.ref = f"A4:T{KB_LAST}"
ws.freeze_panes = "E5"

# =================================================================== NGUON_CONTENT
NC_FIRST, NC_LAST = 5, 304
ws = sheet("NGUON_CONTENT", "NGUỒN CONTENT – HỌC TỪ FACEBOOK, TIKTOK",
           "Dán link video/kênh hay, ghi “Học được gì”. Cột Nền tảng tự nhận. R1 lọc mỗi thứ Sáu: nguồn “Mới” → phân tích → chuyển ý tưởng sang KHO_KICH_BAN. Xem video ngay: tab Nguồn content trong file HTML.",
           [8, 12, 46, 11, 22, 16, 12, 46, 10, 40, 12, 14, 16])
header(ws, 4, ["ID", "Ngày thêm", "Link", "Nền tảng", "Kênh / tác giả", "Loại", "View", "Học được gì (hook, cấu trúc, câu hỏi…)",
               "Áp dụng cho", "Ý tưởng chuyển thể", "ID kịch bản", "Người thêm", "Trạng thái"])
for r in range(NC_FIRST, NC_LAST + 1):
    i = r - NC_FIRST
    s = D.SOURCES[i] if i < len(D.SOURCES) else None
    put(ws, r, 1, s["id"] if s else None, F_INPUT, INPUT_FILL, align=CENTER)
    put(ws, r, 2, dt.date.fromisoformat(s["added"]) if s else None, F_INPUT, INPUT_FILL, "dd/mm/yyyy", CENTER)
    c = put(ws, r, 3, s["url"] if s else None, F_URL if s else F_INPUT, None if s else INPUT_FILL)
    if s:
        c.hyperlink = s["url"]
    put(ws, r, 4, f'=IF(C{r}="","",IF(ISNUMBER(SEARCH("tiktok",C{r})),"TikTok",IF(OR(ISNUMBER(SEARCH("facebook",C{r})),ISNUMBER(SEARCH("fb.watch",C{r}))),"Facebook",IF(OR(ISNUMBER(SEARCH("youtube",C{r})),ISNUMBER(SEARCH("youtu.be",C{r}))),"YouTube",IF(ISNUMBER(SEARCH("instagram",C{r})),"Instagram","Khác")))))', align=CENTER)
    for col, key, fmt in [(5, "author", None), (6, "type", None), (7, "views", "#,##0"), (8, "learned", None), (9, "apply", None),
                          (10, "idea", None), (11, None, None), (12, "by", None), (13, "status", None)]:
        put(ws, r, col, (s.get(key) if (s and key) else None), F_INPUT, INPUT_FILL, fmt, CENTER if col in (7, 9, 11, 13) else WRAP)
    if s:
        ws.row_dimensions[r].height = 45
dv_list(ws, D.SOURCE_TYPES, f"F{NC_FIRST}:F{NC_LAST}")
dv_list(ws, D.PILLAR_CODES + ["Tất cả"], f"I{NC_FIRST}:I{NC_LAST}")
dv_list(ws, D.SOURCE_STATUSES, f"M{NC_FIRST}:M{NC_LAST}")
for st, colr in {"Mới": "FFF59D", "Đã phân tích": "BBDEFB", "Đã thành kịch bản": "C8E6C9", "Bỏ qua": "E0E0E0"}.items():
    ws.conditional_formatting.add(f"M{NC_FIRST}:M{NC_LAST}", CellIsRule(operator="equal", formula=[f'"{st}"'], fill=PatternFill("solid", fgColor=colr)))
ws.auto_filter.ref = f"A4:M{NC_LAST}"
ws.freeze_panes = "D5"

# =================================================================== THEO_DOI_NGAY (1 dòng/ngày)
ws = sheet("THEO_DOI_NGAY", "THEO DÕI SỐ LIỆU HẰNG NGÀY",
           "R4 nhập mỗi sáng trước 10:00 (ô xanh). Ví dụ: 19/10 · Kế hoạch 5 · Đã đăng 5 · View 25.000 · Click 120 · Đơn 3 · Hoa hồng 45.000. Đơn theo từng video xem ở báo cáo Shopee lọc Sub_id.",
           [12, 7, 10, 10, 11, 12, 10, 8, 14, 10, 10, 34])
header(ws, 4, ["Ngày", "Tuần", "Kế hoạch", "Đã đăng", "% hoàn thành", "View", "Click", "Đơn", "Hoa hồng (đ)", "Click/View", "Đơn/Click", "Ghi chú"])
d0 = D.WEEKS[0]["start"]
for k in range(DAYS):
    r = TD_FIRST + k
    put(ws, r, 1, d0 + dt.timedelta(days=k), fmt="dd/mm/yyyy", align=CENTER)
    put(ws, r, 2, f'=IFERROR(INDEX(LO_TRINH!$A$4:$A$15,MATCH(A{r},LO_TRINH!$B$4:$B$15,1)),"")', F_LINK, align=CENTER)
    put(ws, r, 3, f"=IFERROR(INDEX(LO_TRINH!$I$4:$I$15,MATCH(A{r},LO_TRINH!$B$4:$B$15,1)),0)", F_LINK, align=CENTER)
    for col in (4, 6, 7, 8, 9):
        put(ws, r, col, None, F_INPUT, INPUT_FILL, "#,##0", CENTER)
    put(ws, r, 5, f'=IF(OR(C{r}=0,D{r}=""),"",D{r}/C{r})', fmt="0%", align=CENTER)
    put(ws, r, 10, f'=IF(OR(F{r}="",F{r}=0),"",G{r}/F{r})', fmt="0.00%", align=CENTER)
    put(ws, r, 11, f'=IF(OR(G{r}="",G{r}=0),"",H{r}/G{r})', fmt="0.0%", align=CENTER)
    put(ws, r, 12, None, F_INPUT, INPUT_FILL)
ws.conditional_formatting.add(f"E{TD_FIRST}:E{TD_LAST}", FormulaRule(formula=[f'AND(E{TD_FIRST}<>"",E{TD_FIRST}<0.9)'], fill=WARN_FILL))
ws.conditional_formatting.add(f"A{TD_FIRST}:L{TD_LAST}", FormulaRule(formula=[f"$A{TD_FIRST}=TODAY()"], fill=PatternFill("solid", fgColor="E3F2FD")))
ws.auto_filter.ref = f"A4:L{TD_LAST}"
ws.freeze_panes = "C5"

# =================================================================== CONG_VIEC
ws = sheet("CONG_VIEC", "BẢNG CÔNG VIỆC & GIAO VIỆC",
           "Việc thiết lập và mốc lớn. Việc quá hạn chưa xong tự tô đỏ. Thêm việc mới ở dòng trống bên dưới.",
           [7, 66, 30, 18, 12, 13, 14, 30])
header(ws, 4, ["Mã", "Việc", "Vai trò phụ trách", "Người làm", "Hạn", "Trạng thái", "Giai đoạn", "Ghi chú"])
role_name = {ro["code"]: ro["role"] for ro in D.ROLES}
for r in range(CV_FIRST, CV_LAST + 1):
    i = r - CV_FIRST
    t = D.TASKS[i] if i < len(D.TASKS) else None
    put(ws, r, 1, t[0] if t else None, F_BOLD, align=CENTER)
    put(ws, r, 2, t[1] if t else None, F_INPUT if not t else F_BASE, None if t else INPUT_FILL)
    put(ws, r, 3, f"{t[2]} · {role_name[t[2]]}" if t else None, F_INPUT if not t else F_BASE, None if t else INPUT_FILL)
    put(ws, r, 4, None, F_INPUT, INPUT_FILL)
    put(ws, r, 5, t[3] if t else None, F_INPUT, INPUT_FILL, "dd/mm/yyyy", CENTER)
    put(ws, r, 6, "Chưa làm" if t else None, F_INPUT, INPUT_FILL, align=CENTER)
    put(ws, r, 7, t[4] if t else None, align=CENTER)
    put(ws, r, 8, None, F_INPUT, INPUT_FILL)
dv_list(ws, ["Chưa làm", "Đang làm", "Xong", "Bị chặn"], f"F{CV_FIRST}:F{CV_LAST}")
ws.conditional_formatting.add(f"A{CV_FIRST}:H{CV_LAST}", FormulaRule(formula=[f'AND($E{CV_FIRST}<>"",$E{CV_FIRST}<TODAY(),$F{CV_FIRST}<>"Xong")'], fill=WARN_FILL))
ws.conditional_formatting.add(f"F{CV_FIRST}:F{CV_LAST}", CellIsRule(operator="equal", formula=['"Xong"'], fill=PatternFill("solid", fgColor="C8E6C9")))
ws.auto_filter.ref = f"A4:H{CV_LAST}"
ws.freeze_panes = "C5"

# =================================================================== DU_BAO
ws = sheet("DU_BAO", "DỰ BÁO HOA HỒNG THEO LỘ TRÌNH",
           "Ô xanh dương là GIẢ ĐỊNH – thay bằng số thực của page sau 2 tuần (lấy từ THEO_DOI_NGAY). Cột “Thực tế” tự cộng từ số liệu nhập.",
           [36, 14, 14, 12, 10, 16, 16, 12])
section(ws, 3, "Giả định")
rows = [("View trung bình/video", D.FORECAST["views"], F_INPUT, INPUT_FILL, "#,##0", "Giả định – chưa có số thật của page. Thay bằng view TB thực tế (việc T08)."),
        ("Tỷ lệ click link/view", D.FORECAST["ctr"], F_INPUT, INPUT_FILL, "0.00%", "Giả định 0,5% (link ở bình luận). Ngưỡng cảnh báo 0,3% ở KPI_QUY_TAC."),
        ("Tỷ lệ ra đơn/click", D.FORECAST["cr"], F_INPUT, INPUT_FILL, "0.0%", "Giả định 3% cho thực phẩm chức năng, đồ dùng 150–500k."),
        ("Hoa hồng TB/đơn (SP hạng A)", f'=IFERROR(AVERAGEIF(SAN_PHAM!$J${P_FIRST}:$J${P_LAST},"A",SAN_PHAM!$G${P_FIRST}:$G${P_LAST}),AVERAGE(SAN_PHAM!$G${P_FIRST}:$G${P_LAST}))', F_LINK, None, "#,##0", "Tự tính: trung bình HH/đơn của sản phẩm hạng A ở SAN_PHAM (theo Shopee).")]
for i, (k, v, f, fl, fmt, note) in enumerate(rows):
    r = 4 + i
    put(ws, r, 1, k, F_BOLD); put(ws, r, 2, v, f, fl, fmt, CENTER); put(ws, r, 3, note, F_SUB)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
header(ws, 10, ["Tuần", "Video/tuần", "View dự kiến", "Click", "Đơn", "Hoa hồng dự kiến (đ)", "Hoa hồng thực tế (đ)", "Thực tế/Dự kiến"])
for i, w in enumerate(D.WEEKS):
    r = 11 + i
    put(ws, r, 1, f'="Tuần "&LO_TRINH!A{4+i}&" ("&TEXT(LO_TRINH!B{4+i},"dd/mm")&")"', F_LINK)
    put(ws, r, 2, f"=LO_TRINH!J{4+i}", F_LINK, fmt="#,##0", align=CENTER)
    put(ws, r, 3, f"=B{r}*$B$4", fmt="#,##0", align=CENTER)
    put(ws, r, 4, f"=C{r}*$B$5", fmt="#,##0", align=CENTER)
    put(ws, r, 5, f"=D{r}*$B$6", fmt="#,##0", align=CENTER)
    put(ws, r, 6, f"=E{r}*$B$7", fmt="#,##0", align=CENTER)
    put(ws, r, 7, f"=SUMIFS(THEO_DOI_NGAY!$I${TD_FIRST}:$I${TD_LAST},THEO_DOI_NGAY!$B${TD_FIRST}:$B${TD_LAST},LO_TRINH!A{4+i})", F_LINK, fmt="#,##0", align=CENTER)
    put(ws, r, 8, f'=IF(F{r}=0,"",G{r}/F{r})', fmt="0%", align=CENTER)
r = 11 + len(D.WEEKS)
put(ws, r, 1, "TỔNG 12 TUẦN", F_BOLD, SUB_FILL)
for col in range(2, 8):
    L = get_column_letter(col)
    put(ws, r, col, f"=SUM({L}11:{L}{r-1})", F_BOLD, SUB_FILL, "#,##0", CENTER)
put(ws, r, 8, f'=IF(F{r}=0,"",G{r}/F{r})', F_BOLD, SUB_FILL, "0%", CENTER)
put(ws, r + 1, 1, f"Hoa hồng dự kiến/tháng khi đủ {MAXD} video/ngày", F_BOLD)
put(ws, r + 1, 6, f"=F{11+6}/7*30", F_BOLD, fmt="#,##0", align=CENTER)
m = r + 1
section(ws, m + 2, f"Bài toán hòa vốn (khi đủ {MAXD} video/ngày)")
be = [("Số người khi đủ công suất", f"=NHAN_SU!B{11+len(D.ROLES)}", F_LINK, None, "0", "Lấy từ NHAN_SU. Giai đoạn 1 chỉ cần 2 người kiêm nhiệm."),
      ("Chi phí TB/người/tháng (đ)", D.FORECAST["cost"], F_INPUT, INPUT_FILL, "#,##0", "Giả định 8 triệu/người – điền mức thực tế."),
      ("Chi phí nhân sự/tháng (đ)", f"=B{m+3}*B{m+4}", F_BASE, None, "#,##0", ""),
      ("Lãi/lỗ mỗi tháng (đ)", f"=F{m}-B{m+5}", F_BOLD, None, "#,##0;[Red]-#,##0", "Hoa hồng/tháng trừ chi phí nhân sự."),
      ("View TB/video cần để hòa vốn", f'=IF($B$5*$B$6*$B$7=0,"",B{m+5}/(MAX(LO_TRINH!I4:I15)*30*$B$5*$B$6*$B$7))', F_BOLD, None, "#,##0", "Mốc quyết định tăng người: chỉ tăng khi view thực tế vượt mốc này.")]
for i, (k, v, f, fl, fmt, note) in enumerate(be):
    rr = m + 3 + i
    put(ws, rr, 1, k, F_BOLD); put(ws, rr, 2, v, f, fl, fmt, CENTER); put(ws, rr, 3, note, F_SUB)
    ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=8)
ws.conditional_formatting.add(f"B{m+6}", CellIsRule(operator="lessThan", formula=["0"], fill=WARN_FILL))

# =================================================================== KPI_QUY_TAC
ws = sheet("KPI_QUY_TAC", "KPI & QUY TẮC CẮT / NHÂN", "Áp dụng khi review Chủ nhật.", [32, 22, 22, 60])
header(ws, 4, ["Chỉ số", "Theo dõi ở", "Ngưỡng", "Hành động"])
for i, k in enumerate(D.KPI_RULES):
    for col, v in enumerate(k, 1):
        put(ws, 5 + i, col, v, F_BOLD if col == 1 else F_BASE)
r = 5 + len(D.KPI_RULES) + 1
put(ws, r, 1, "Mỗi Chủ nhật: lấy 3 video ra đơn nhiều nhất (báo cáo Shopee lọc Sub_id) → làm 5 biến thể cho tuần sau.", F_BOLD)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)

# =================================================================== RUI_RO
ws = sheet("RUI_RO", "SỔ RỦI RO", "Rà lại mỗi tháng.", [56, 12, 76, 14, 16])
header(ws, 4, ["Rủi ro", "Mức", "Cách phòng tránh", "Chủ quản", "Tình trạng"])
lvl = {"Cao": WARN_FILL, "Trung bình": PatternFill("solid", fgColor="FFF59D"), "Thấp": PatternFill("solid", fgColor="E8F5E9")}
for i, k in enumerate(D.RISKS):
    r = 5 + i
    put(ws, r, 1, k[0]); put(ws, r, 2, k[1], F_BOLD, lvl[k[1]], align=CENTER); put(ws, r, 3, k[2]); put(ws, r, 4, k[3], align=CENTER)
    put(ws, r, 5, "Đang theo dõi", F_INPUT, INPUT_FILL, align=CENTER)
    ws.row_dimensions[r].height = 32
dv_list(ws, ["Đang theo dõi", "Đã xử lý", "Đã xảy ra"], f"E5:E{4+len(D.RISKS)}")

# =================================================================== TAI_NGUYEN
ws = sheet("TAI_NGUYEN", "TÀI NGUYÊN & CÔNG CỤ", "Đường dẫn trên máy và skill dùng trong Claude Code.", [40, 80, 46])
header(ws, 4, ["Tài nguyên", "Đường dẫn / tên", "Dùng để"])
for i, k in enumerate(D.RESOURCES):
    r = 5 + i
    put(ws, r, 1, k[0], F_BOLD)
    c = put(ws, r, 2, k[1], F_URL if (":\\" in k[1] or k[1].startswith("http")) else F_BASE)
    if ":\\" in k[1]:
        c.hyperlink = "file:///" + k[1].replace("\\", "/")
    elif k[1].startswith("http"):
        c.hyperlink = k[1]
    put(ws, r, 3, k[2])

# =================================================================== TONG_QUAN
ws = sheet("TONG_QUAN", "TỔNG QUAN ĐIỀU HÀNH", "Mở mỗi sáng. Mọi số đều tự tính từ các sheet khác – không nhập tay ở đây.",
           [36, 16, 14, 14, 14, 14, 14, 16])
td = lambda col: f"THEO_DOI_NGAY!${col}${TD_FIRST}:${col}${TD_LAST}"
kb = lambda col: f"KHO_KICH_BAN!${col}${KB_FIRST}:${col}${KB_LAST}"
section(ws, 3, "Chỉ số chính")
kpis = [
    ("Video kế hoạch đến hôm nay", f'=SUMIFS({td("C")},{td("A")},"<="&TODAY())', "#,##0"),
    ("Video đã đăng", f"=SUM({td('D')})", "#,##0"),
    ("% hoàn thành kế hoạch", '=IF(B4=0,"",B5/B4)', "0%"),
    ("Tổng view", f"=SUM({td('F')})", "#,##0"),
    ("Tổng click", f"=SUM({td('G')})", "#,##0"),
    ("Tổng đơn", f"=SUM({td('H')})", "#,##0"),
    ("Tổng hoa hồng (đ)", f"=SUM({td('I')})", "#,##0"),
    ("Hoa hồng/video đã đăng (đ)", '=IF(B5=0,"",B10/B5)', "#,##0"),
    ("Hoa hồng 7 ngày gần nhất (đ)", f'=SUMIFS({td("I")},{td("A")},">"&(TODAY()-7),{td("A")},"<="&TODAY())', "#,##0"),
]
for i, (k, f, fmt) in enumerate(kpis):
    r = 4 + i
    put(ws, r, 1, k, F_BOLD); put(ws, r, 2, f, F_LINK if "THEO" in f else F_BASE, fmt=fmt, align=CENTER)
section(ws, 15, "Theo trụ cột (từ KHO_KICH_BAN, video “Đã đăng”)")
header(ws, 16, ["Trụ cột", "Video/ngày hiện tại", "Đã đăng", "View", "Click", "Đơn", "Click/View", "Đơn/video"])
for i, pl in enumerate(D.PILLARS):
    r = 17 + i
    L = get_column_letter(4 + i)
    put(ws, r, 1, f"{pl['code']} · {pl['name']}", F_BOLD)
    put(ws, r, 2, f'=IFERROR(INDEX(LO_TRINH!{L}$4:{L}$15,MATCH(TODAY(),LO_TRINH!$B$4:$B$15,1)),0)', F_LINK, align=CENTER)
    put(ws, r, 3, f'=COUNTIFS({kb("B")},"{pl["code"]}",{kb("M")},"Đã đăng")', F_LINK, fmt="#,##0", align=CENTER)
    for col, src in [(4, "Q"), (5, "R"), (6, "S")]:
        put(ws, r, col, f'=SUMIFS({kb(src)},{kb("B")},"{pl["code"]}",{kb("M")},"Đã đăng")', F_LINK, fmt="#,##0", align=CENTER)
    put(ws, r, 7, f'=IF(D{r}=0,"",E{r}/D{r})', fmt="0.00%", align=CENTER)
    put(ws, r, 8, f'=IF(C{r}=0,"",F{r}/C{r})', fmt="0.00", align=CENTER)
rt = 17 + NPIL
put(ws, rt, 1, "Tổng", F_BOLD, SUB_FILL)
for col in range(2, 7):
    L = get_column_letter(col)
    put(ws, rt, col, f"=SUM({L}17:{L}{rt-1})", F_BOLD, SUB_FILL, "#,##0", CENTER)
put(ws, rt, 7, f'=IF(D{rt}=0,"",E{rt}/D{rt})', F_BOLD, SUB_FILL, "0.00%", CENTER)
put(ws, rt, 8, f'=IF(C{rt}=0,"",F{rt}/C{rt})', F_BOLD, SUB_FILL, "0.00", CENTER)
r = rt + 2
section(ws, r, "Theo loại video"); r += 1
header(ws, r, ["Loại video", "Đã đăng", "View TB", "Đơn", "Đơn/video"])
for t in D.VIDEO_TYPES:
    r += 1
    put(ws, r, 1, t, F_BOLD)
    put(ws, r, 2, f'=COUNTIFS({kb("C")},A{r},{kb("M")},"Đã đăng")', F_LINK, fmt="#,##0", align=CENTER)
    put(ws, r, 3, f'=IF(B{r}=0,"",SUMIFS({kb("Q")},{kb("C")},A{r},{kb("M")},"Đã đăng")/B{r})', F_LINK, fmt="#,##0", align=CENTER)
    put(ws, r, 4, f'=SUMIFS({kb("S")},{kb("C")},A{r},{kb("M")},"Đã đăng")', F_LINK, fmt="#,##0", align=CENTER)
    put(ws, r, 5, f'=IF(B{r}=0,"",D{r}/B{r})', fmt="0.00", align=CENTER)
r += 2
section(ws, r, "Dây chuyền kịch bản (KHO_KICH_BAN)"); r += 1
header(ws, r, ["Trạng thái", "Số video"])
for s in D.STATUSES:
    r += 1
    put(ws, r, 1, s); put(ws, r, 2, f'=COUNTIF({kb("M")},A{r})', F_LINK, align=CENTER)
r += 2
section(ws, r, "Việc cần chú ý")
alerts = [
    ("Link sản phẩm chưa kiểm tra", f'=COUNTIF(SAN_PHAM!$Q${P_FIRST}:$Q${P_LAST},"Chưa kiểm tra")'),
    ("Link sản phẩm “Cần sửa”", f'=COUNTIF(SAN_PHAM!$Q${P_FIRST}:$Q${P_LAST},"Cần sửa")'),
    ("Công việc quá hạn chưa xong", f'=COUNTIFS(CONG_VIEC!$E${CV_FIRST}:$E${CV_LAST},"<"&TODAY(),CONG_VIEC!$F${CV_FIRST}:$F${CV_LAST},"<>Xong")'),
    ("Công việc đã xong", f'=COUNTIF(CONG_VIEC!$F${CV_FIRST}:$F${CV_LAST},"Xong")'),
    ("Prompt đang chờ duyệt", f'=COUNTIF({kb("M")},"Chờ duyệt")'),
    ("Nguồn content mới chưa phân tích", f'=COUNTIF(NGUON_CONTENT!$M$5:$M$304,"Mới")'),
    (f"Nhân sự cần khi đủ {MAXD} video/ngày", f"=NHAN_SU!B{11+len(D.ROLES)}"),
    ("Tổng video theo lộ trình 12 tuần", "=LO_TRINH!J16"),
    ("Hoa hồng dự kiến 12 tuần (đ, theo giả định)", f"=DU_BAO!F{11+len(D.WEEKS)}"),
]
for i, (k, f) in enumerate(alerts):
    rr = r + 1 + i
    put(ws, rr, 1, k, F_BOLD); put(ws, rr, 2, f, F_LINK, fmt="#,##0", align=CENTER)
ws.conditional_formatting.add(f"B{r+1}:B{r+3}", CellIsRule(operator="greaterThan", formula=["0"], fill=WARN_FILL))

# thứ tự sheet
order = ["HUONG_DAN", "TONG_QUAN", "KENH", "LO_TRINH", "NHAN_SU", "QUY_TRINH", "LICH_DANG", "SAN_PHAM",
         "KHO_KICH_BAN", "NGUON_CONTENT", "THEO_DOI_NGAY", "CONG_VIEC", "DU_BAO", "KPI_QUY_TAC", "RUI_RO", "TAI_NGUYEN"]
wb._sheets = [wb[n] for n in order]
tab_colors = {"NGUON_CONTENT": "7C3AED", "TONG_QUAN": "0F4C5C", "KHO_KICH_BAN": "DB2777", "THEO_DOI_NGAY": "16A34A", "CONG_VIEC": "EA580C", "SAN_PHAM": "2563EB"}
for n, c in tab_colors.items():
    wb[n].sheet_properties.tabColor = c
wb.active = 1
wb.calculation.fullCalcOnLoad = True


def build(path):
    wb.save(path)
    return path
