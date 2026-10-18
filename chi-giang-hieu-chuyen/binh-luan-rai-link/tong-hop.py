"""Gộp sản phẩm kênh Chị Giang Hiểu Chuyện: masterplan (san-pham-kenh-chi-giang.xlsx) + đàn ông 40+ (research/dan-ong-40)
thành 1 file Excel: 4 bình luận x 10 link, xếp từ đúng ngách đến rộng, kèm sheet đoạn chữ copy-dán."""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(r"D:\video-ai-aff")
HERE = Path(__file__).parent
SUB = "ChiGiang-Tap01-FBReels"

src = {}
for r in openpyxl.load_workbook(ROOT / "san-pham-kenh-chi-giang.xlsx").active.iter_rows(min_row=2, values_only=True):
    if r[0]:
        src[f"M{r[0]}"] = dict(title=r[3], price=r[4], comm=r[5], sold="", link=r[6], src="Masterplan kênh")
for r in openpyxl.load_workbook(ROOT / "research/dan-ong-40/san-pham-dan-ong-40.xlsx").active.iter_rows(min_row=2, values_only=True):
    if r[0]:
        src[f"D{r[0]}"] = dict(title=r[3], price=r[4], comm=r[5], sold=r[6], link=r[8], src="Đàn ông 40+ (KaloData)")

GROUPS = [
    ("Phong độ & tiết niệu", "💪", "Đúng ngách", [
        ("M1", "Kẽm DSD 150mg"), ("M2", "Sâm nhung bổ thận DSD"), ("M3", "Vương Bảo tiền liệt tuyến"), ("M4", "Viên thận Botamax"),
        ("M5", "Kẽm ADA 150mg"), ("M6", "Bọt vệ sinh nam Oniiz"), ("M8", "Bao cao su Durex Performa"), ("M18", "Quần boxer thông hơi"),
        ("D10", "Xịt thơm miệng"), ("M23", "Trà mâm xôi (quà cho vợ)")]),
    ("Huyết áp, xương khớp, tiểu đường", "🩺", "Đúng ngách", [
        ("M11", "Máy đo huyết áp Omron Ez"), ("D18", "Máy đo huyết áp Omron HEM-7156T"), ("M12", "Máy đo huyết áp giọng nói"),
        ("M13", "Dây thìa canh Diabetna"), ("M9", "Gel xương khớp Phục Thể An"), ("M10", "Thiên Ma thấu cốt hoàn"),
        ("M15", "Magie chống chuột rút đêm"), ("M14", "Miếng dán chân ngủ ngon"), ("D9", "Kính lão gấp gọn"), ("D4", "Đai chống gù lưng")]),
    ("Trẻ lại & thư giãn", "😎", "Mở rộng", [
        ("M16", "Dầu gội phủ bạc Nezha"), ("M17", "Dầu gội nhuộm nâu Korolo"), ("M19", "Đai lưng tự nóng Kazumi"),
        ("M20", "Gối massage Kingtech"), ("M21", "Súng massage cầm tay"), ("M22", "Kềm cắt móng Kềm Nghĩa"),
        ("D8", "Xịt vệ sinh giày"), ("D15", "Ly giữ nhiệt 84Hydro"), ("D6", "Loa karaoke 2 mic"), ("D17", "Vợt muỗi kiêm đèn bắt muỗi")]),
    ("Nhà cửa, xe cộ, đồ nghề", "🔧", "Mở rộng", [
        ("D2", "Xịt phủ kính nano"), ("D5", "Kích bình bơm lốp ô tô"), ("D20", "Máy rửa xe cao áp"), ("D16", "Thẻ định vị Ugreen"),
        ("D13", "Camera xoay 360"), ("D14", "Đèn năng lượng mặt trời"), ("D12", "Đèn pin đội đầu"), ("D19", "Máy khoan pin"),
        ("D7", "Nước thông cống"), ("D3", "Keo bọt nở chống thấm")]),
]

thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Sản phẩm theo nhóm"
hdr = ["Bình luận", "Số", "Tên ngắn", "Tên sản phẩm Shopee", "Giá", "Hoa hồng", "Lượt bán", "Link Shopee", "Link kèm Sub_id (dán vào bình luận)", "Nguồn"]
ws.append(hdr)
for c in ws[1]:
    c.fill = PatternFill("solid", start_color="1F4E78")
    c.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws2 = wb.create_sheet("Bình luận copy-dán")
ws2.append(["Bình luận", "Đoạn chữ (copy nguyên ô, dán vào bình luận)"])
for c in ws2[1]:
    c.fill = PatternFill("solid", start_color="1F4E78")
    c.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")

used = set()
for n, (title, emo, tier, items) in enumerate(GROUPS, 1):
    ws.append([f"BL{n:02d} · {emo} {title.upper()}  ({tier})"])
    rn = ws.max_row
    ws.merge_cells(start_row=rn, start_column=1, end_row=rn, end_column=len(hdr))
    niche = tier == "Đúng ngách"
    ws.cell(rn, 1).fill = PatternFill("solid", start_color="E3F2FD" if niche else "FFF3E0")
    ws.cell(rn, 1).font = Font(name="Arial", size=12, bold=True, color="0D47A1" if niche else "E65100")
    lines = ["Dạ em cám ơn anh chị đã xem hết video ạ 🙏 Đồ anh em hay hỏi, em để link bên dưới:", "",
             f"{emo} {title.upper()} 👇"]
    for i, (k, short) in enumerate(items, 1):
        p = src[k]
        link = p["link"].strip()
        assert link not in used, f"link trùng {k}"
        used.add(link)
        sub = f"{link}?sub_id={SUB}-{link.rstrip('/').split('/')[-1]}"
        lines.append(f"{i}. {short} 👉 {sub}")
        ws.append([f"BL{n:02d}", i, short, p["title"], p["price"], p["comm"], p["sold"], link, sub, p["src"]])
        rn = ws.max_row
        for ci in range(1, len(hdr) + 1):
            c = ws.cell(rn, ci)
            c.font = Font(name="Arial", size=10)
            c.border = border
            c.alignment = Alignment(horizontal="center" if ci in (1, 2, 5, 6, 7) else "left", vertical="center", wrap_text=True)
        for ci in (8, 9):
            ws.cell(rn, ci).hyperlink = ws.cell(rn, ci).value
            ws.cell(rn, ci).font = Font(name="Arial", size=10, color="0563C1", underline="single")
    ws2.append([f"BL{n:02d} · {title} ({tier})", "\n".join(lines)])
    ws2.cell(ws2.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws2.cell(ws2.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")

for i, w in enumerate([10, 5, 28, 50, 13, 10, 10, 30, 60, 20], 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "C2"
ws2.column_dimensions["A"].width = 30
ws2.column_dimensions["B"].width = 110

out = HERE / "san-pham-chi-giang-tong-hop.xlsx"
wb.save(out)
print("Đã lưu", out, "|", len(used), "link")
